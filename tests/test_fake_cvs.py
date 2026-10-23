"""Edge cases found by stress-testing with fictional CVs and job descriptions."""
import io

import docx
import pytest

from src.matcher import match
from src.parsing import extract_text
from src.profile import cv_years, required_years


@pytest.mark.parametrize("cv,years", [
    ("Experience\nFrontend Developer, 2021-03 to 2024-06\nIntern, 06/2020 - 08/2020", 3.6),
    ("EXPERIENCE\nData Scientist, Sept '19 - Present\nML Engineer, Jun. 2017 to Aug. 2019", 9.3),
    # No experience heading: education dates must not count as work.
    ("Ali\nEDUCATION\nBSc Computer Science, 2022 - 2026\nPROJECTS\nChatbot, 2025\nSKILLS\nPython", 0.0),
])
def test_cv_date_edge_cases(cv, years):
    assert cv_years(cv, today=2026, today_month=10) == pytest.approx(years)


@pytest.mark.parametrize("jd,years", [
    ("Frontend Developer\n3-5 yrs with React. Company founded 20 years ago.", 3),
    ("DevOps Engineer\nMinimum three years of experience with AWS.", 3),
    ("Maths Tutor\nNo experience required.", 0),
])
def test_required_years_edge_cases(jd, years):
    assert required_years(jd) == years


def test_docx_table_stays_in_its_section():
    d = docx.Document()
    d.add_paragraph("EXPERIENCE")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Software Engineer", "Feb 2019 – Present"
    d.add_paragraph("SKILLS")
    d.add_paragraph("Python")
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text(buf.getvalue(), "cv.docx")
    assert text.index("Feb 2019") < text.index("SKILLS")
    assert cv_years(text, today=2026, today_month=10) == pytest.approx(7.7)


def test_no_penalty_when_jd_states_no_level():
    r = match("Ahmed\nSKILLS\nPython", "Python Developer\nWrite Python scripts.", use_embeddings=False)
    assert r.components["experience"] == 100.0
